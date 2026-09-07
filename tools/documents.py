"""
tools/documents.py

Generates real, downloadable artifacts:
  - generate_invoice_pdf: a GST-correct invoice (reportlab), only for
    an already-FINALIZED bill (invoicing a draft doesn't make sense --
    nothing's actually been sold yet).
  - generate_analysis_deck: a PPTX with a real, editable native chart
    (python-pptx's CategoryChartData + add_chart), not a pasted image.

Files are written to generated/ and the path is returned in the
result dict -- bot/agent.py picks that up and bot/main.py sends it
as an actual Telegram document (see pop_last_generated_file there).
"""

import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from tools.billing import get_bill_draft
from tools.preferences import get_all_preferences
from tools.analytics import get_sales_summary

GENERATED_DIR = "generated"
os.makedirs(GENERATED_DIR, exist_ok=True)


def generate_invoice_pdf(bill_id: int) -> dict:
    bill = get_bill_draft(bill_id)
    if not bill["ok"]:
        return bill
    if bill["status"] != "finalized":
        return {
            "ok": False,
            "error": "not_finalized",
            "detail": "Only finalized bills can be invoiced -- finalize it first.",
        }

    prefs = get_all_preferences()["preferences"]
    shop_name = prefs.get("shop_name", "Kirana Store")
    shop_gstin = prefs.get("shop_gstin", "")

    path = os.path.join(GENERATED_DIR, f"invoice_{bill_id}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"<b>{shop_name}</b>", styles["Title"]),
    ]
    if shop_gstin:
        elements.append(Paragraph(f"GSTIN: {shop_gstin}", styles["Normal"]))
    elements.append(Paragraph(f"Invoice for Bill #{bill_id}", styles["Heading2"]))
    elements.append(Spacer(1, 10))

    table_data = [["Item", "HSN", "Qty", "Rate", "Taxable Value", "CGST", "SGST", "Total"]]
    for item in bill["items"]:
        taxable = item["line_total"] - item["cgst_amt"] - item["sgst_amt"]
        table_data.append([
            item["name"],
            item["hsn_code"],
            f'{item["qty"]:g}',
            f'{item["unit_price"]:.2f}',
            f'{taxable:.2f}',
            f'{item["cgst_amt"]:.2f}',
            f'{item["sgst_amt"]:.2f}',
            f'{item["line_total"]:.2f}',
        ])
    table_data.append(["", "", "", "", "", "", "Total", f'{bill["total"]:.2f}'])

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -2), 0.5, colors.grey),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    elements.append(table)
    doc.build(elements)

    return {"ok": True, "bill_id": bill_id, "file_path": path}


def generate_analysis_deck(date_from: str = None, date_to: str = None) -> dict:
    summary = get_sales_summary(date_from, date_to)
    if not summary["ok"]:
        return summary

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Sales Analysis"
    slide.placeholders[1].text = f'{summary["date_from"]} to {summary["date_to"]}'

    # Overview slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Overview"
    body = slide.placeholders[1].text_frame
    body.text = f'Total Sales: Rs {summary["total_sales"]:.2f}'
    body.add_paragraph().text = f'Tax Collected: Rs {summary["total_tax_collected"]:.2f}'
    for mode, amt in summary["by_payment_mode"].items():
        body.add_paragraph().text = f'{mode or "unspecified"}: Rs {amt:.2f}'

    # Top items -- a REAL native chart, not a pasted-in image
    if summary["top_items"]:
        chart_data = CategoryChartData()
        chart_data.categories = [i["name"] for i in summary["top_items"]]
        chart_data.add_series("Qty Sold", [i["qty_sold"] for i in summary["top_items"]])

        slide = prs.slides.add_slide(blank_layout)
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        title_box.text_frame.text = "Top Items by Quantity Sold"
        title_box.text_frame.paragraphs[0].font.size = Pt(28)
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.2), Inches(9), Inches(5), chart_data
        )

    path = os.path.join(GENERATED_DIR, f'analysis_{summary["date_from"]}_{summary["date_to"]}.pptx')
    prs.save(path)

    return {"ok": True, "file_path": path, **summary}